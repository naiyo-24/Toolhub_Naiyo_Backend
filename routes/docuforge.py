from fastapi import APIRouter, HTTPException, File, UploadFile, Form, Response
import mimetypes
from schemas.docuforge import ResumeBuilderRequest, IDCardRequest
from routes.ai_tools import generate_ai_response, upload_to_gemini
from routes.student_tools import clean_json_string
import json

router = APIRouter()

@router.post("/resume-builder")
def resume_builder(req: ResumeBuilderRequest):
    prompt = f"Generate a professional ATS-friendly resume for {req.full_name} ({req.email}, {req.phone}). Target role: {req.target_role}. Education: {req.education}. Experience: {req.experience}. Skills: {', '.join(req.skills)}. Return a JSON object with keys: 'summary', 'experience_bullets', 'education_bullets', 'skills'."
    raw_json = clean_json_string(generate_ai_response(prompt, json_mode=True))
    try:
        return {"resume": json.loads(raw_json)}
    except json.JSONDecodeError:
        return {"resume": raw_json}

import re
from pypdf import PdfReader
from io import BytesIO

@router.post("/ats-checker")
async def ats_checker(resume: UploadFile = File(...), job_description: str = Form(...)):
    try:
        reader = PdfReader(BytesIO(await resume.read()))
        resume_text = " ".join([page.extract_text() or "" for page in reader.pages]).lower()
        
        # Heuristic to fix PDF kerning issues (where letters are separated by spaces)
        # e.g., "h t m l   d e v e l o p e r"
        words_split = resume_text.split()
        if words_split:
            single_char_count = sum(1 for w in words_split if len(w) == 1 and w.isalpha())
            if single_char_count / len(words_split) > 0.4:
                # Replace double spaces or newlines with a placeholder
                resume_text = resume_text.replace("  ", " _W_ ").replace("\n", " _W_ ")
                # Remove single spaces
                resume_text = resume_text.replace(" ", "")
                # Restore boundaries
                resume_text = resume_text.replace("_W_", " ").replace(",", ", ")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to read PDF: {str(e)}")

    jd_text = job_description.lower()

    def get_words(text):
        words = re.findall(r'\b[a-z]{3,}\b', text)
        stop_words = {"the", "and", "for", "with", "from", "that", "this", "are", "you", "your", "will", "can", "have", "has", "not", "but", "our", "all", "any", "how"}
        return set([w for w in words if w not in stop_words])

    resume_words = get_words(resume_text)
    jd_words = get_words(jd_text)

    if len(jd_words) == 0:
        jd_words = {"experience"} # fallback

    matched = jd_words.intersection(resume_words)
    missing = jd_words.difference(resume_words)
    
    # JDs are full of fluff. Matching 25% of unique non-stop words is practically a perfect match.
    effective_total = max(1, int(len(jd_words) * 0.25))
    score = int((len(matched) / effective_total) * 100)
    score = min(100, score)
    
    formatting_issues = []
    improvement_tips = []
    
    word_count = len(resume_text.split())
    if word_count < 250:
        formatting_issues.append("Resume is too short (under 250 words).")
        improvement_tips.append("Add more details to your bullet points to show impact.")
    elif word_count > 1000:
        formatting_issues.append("Resume is a bit long (over 1000 words).")
        improvement_tips.append("Consider trimming older or less relevant experience.")

    if "education" not in resume_text:
        formatting_issues.append("Could not find an 'Education' section.")
        improvement_tips.append("Add a clearly labeled 'Education' section.")
        
    if "experience" not in resume_text and "employment" not in resume_text and "work history" not in resume_text:
        formatting_issues.append("Could not find an 'Experience' section.")
        improvement_tips.append("Add a clearly labeled 'Experience' or 'Employment' section.")

    missing_list = list(missing)[:12]

    return {
        "ats_analysis": {
            "ats_score": score,
            "missing_keywords": missing_list,
            "formatting_issues": formatting_issues,
            "improvement_tips": improvement_tips
        }
    }

@router.post("/cover-letter")
async def cover_letter(resume: UploadFile = File(...), job_description: str = Form(...), tone: str = Form("Professional")):
    resume_bytes = await resume.read()
    mime = mimetypes.guess_type(resume.filename)[0] or "application/pdf"
    file_uri = upload_to_gemini(resume_bytes, resume.filename, mime)
    
    prompt = f"Write a highly persuasive {tone} cover letter for the following job description:\n{job_description}\nBase the cover letter on the candidate's uploaded resume."
    return {"cover_letter": generate_ai_response(prompt, file_uri=file_uri)}

@router.post("/ocr-scanner")
async def ocr_scanner(image: UploadFile = File(...)):
    image_bytes = await image.read()
    mime = mimetypes.guess_type(image.filename)[0] or "image/jpeg"
    # file_uri = upload_to_gemini(image_bytes, image.filename, mime)
    # prompt = "Extract all the text from this image precisely as it appears. Preserve the formatting as best as possible."
    # return {"extracted_text": generate_ai_response(prompt, file_uri=file_uri)}
    return {"extracted_text": "OCR processing via Google ML Kit should be handled directly on the frontend (client-side). Backend OCR is disabled as per user request."}

from fastapi.responses import StreamingResponse
from pypdf import PdfReader, PdfWriter
from io import BytesIO
from PIL import Image, ImageEnhance
from typing import List

@router.post("/merge-pdf")
async def merge_pdf(files: List[UploadFile] = File(...)):
    merger = PdfWriter()
    for file in files:
        try:
            merger.append(BytesIO(await file.read()))
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Invalid PDF file '{file.filename}': {str(e)}")
    
    out_stream = BytesIO()
    merger.write(out_stream)
    out_stream.seek(0)
    return StreamingResponse(out_stream, media_type="application/pdf", headers={"Content-Disposition": "attachment; filename=merged.pdf"})

@router.post("/split-pdf")
async def split_pdf(file: UploadFile = File(...), page_numbers: str = Form(...)):
    reader = PdfReader(BytesIO(await file.read()))
    writer = PdfWriter()
    
    pages = [int(p.strip()) - 1 for p in page_numbers.split(",") if p.strip().isdigit()]
    for p in pages:
        if 0 <= p < len(reader.pages):
            writer.add_page(reader.pages[p])
            
    out_stream = BytesIO()
    writer.write(out_stream)
    out_stream.seek(0)
    return StreamingResponse(out_stream, media_type="application/pdf", headers={"Content-Disposition": "attachment; filename=split.pdf"})

@router.post("/compress-pdf")
async def compress_pdf(
    file: UploadFile = File(...), 
    quality: int = Form(50), 
    target_size_kb: int = Form(None)
):
    try:
        content = await file.read()
        
        if target_size_kb is not None and target_size_kb > 0:
            target_bytes = target_size_kb * 1024
            current_quality = 80
            best_out_stream = None
            
            while current_quality >= 10:
                reader = PdfReader(BytesIO(content))
                writer = PdfWriter()
                for page in reader.pages:
                    writer.add_page(page)
                
                if reader.metadata is not None:
                    writer.add_metadata(reader.metadata)
                    
                for page in writer.pages:
                    page.compress_content_streams()
                    if hasattr(page, "images"):
                        for img in page.images:
                            try:
                                img.replace(img.image, quality=current_quality)
                            except Exception:
                                pass
                
                out_stream = BytesIO()
                writer.write(out_stream)
                
                if out_stream.getbuffer().nbytes <= target_bytes:
                    best_out_stream = out_stream
                    break
                
                # If quality is 10 and we still haven't met target, just keep it
                if current_quality == 10:
                    best_out_stream = out_stream
                    
                current_quality -= 10
                
            best_out_stream.seek(0)
            return StreamingResponse(best_out_stream, media_type="application/pdf", headers={"Content-Disposition": "attachment; filename=compressed.pdf"})
            
        else:
            reader = PdfReader(BytesIO(content))
            writer = PdfWriter()
            for page in reader.pages:
                writer.add_page(page)
            
            if reader.metadata is not None:
                writer.add_metadata(reader.metadata)
                
            for page in writer.pages:
                page.compress_content_streams()
                
                # Compress images to actually reduce file size
                if hasattr(page, "images"):
                    for img in page.images:
                        try:
                            img.replace(img.image, quality=quality)
                        except Exception:
                            pass
                
            out_stream = BytesIO()
            writer.write(out_stream)
            out_stream.seek(0)
            return StreamingResponse(out_stream, media_type="application/pdf", headers={"Content-Disposition": "attachment; filename=compressed.pdf"})
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to compress PDF: {str(e)}")

@router.post("/document-scan")
async def document_scan(
    images: List[UploadFile] = File(...),
    scan_type: str = Form("magic_color")
):
    enhanced_imgs = []
    
    for img_file in images:
        img_bytes = await img_file.read()
        img = Image.open(BytesIO(img_bytes)).convert("RGB")
        
        if scan_type == "magic_color":
            enhancer = ImageEnhance.Contrast(img)
            img = enhancer.enhance(1.8)
            enhancer = ImageEnhance.Brightness(img)
            img = enhancer.enhance(1.1)
            enhancer = ImageEnhance.Color(img)
            img = enhancer.enhance(1.5)
        elif scan_type == "black_white":
            # Convert to Grayscale then threshold
            img = img.convert("L")
            img = img.point(lambda x: 0 if x < 128 else 255, '1')
            # Convert back to RGB for consistency in saving
            img = img.convert("RGB")
        elif scan_type == "grayscale":
            img = img.convert("L")
            img = img.convert("RGB")
        # original does nothing
        
        enhanced_imgs.append(img)
        
    out_stream = BytesIO()
    
    if len(enhanced_imgs) == 1:
        enhanced_imgs[0].save(out_stream, format="PDF")
    else:
        enhanced_imgs[0].save(out_stream, format="PDF", save_all=True, append_images=enhanced_imgs[1:])
        
    out_stream.seek(0)
    return Response(content=out_stream.getvalue(), media_type="application/pdf", headers={"Content-Disposition": "attachment; filename=scanned_document.pdf"})

@router.post("/image-to-pdf")
async def image_to_pdf(images: List[UploadFile] = File(...)):
    image_objs = []
    for img_file in images:
        img_bytes = await img_file.read()
        img = Image.open(BytesIO(img_bytes)).convert("RGB")
        image_objs.append(img)
        
    if not image_objs:
        raise HTTPException(status_code=400, detail="No images provided")
        
    out_stream = BytesIO()
    image_objs[0].save(out_stream, format="PDF", save_all=True, append_images=image_objs[1:])
    out_stream.seek(0)
    return StreamingResponse(out_stream, media_type="application/pdf", headers={"Content-Disposition": "attachment; filename=converted.pdf"})

import subprocess
import tempfile
import os
import uuid

def convert_with_libreoffice(input_path: str, output_dir: str):
    # Mac LibreOffice binary path
    libreoffice_path = "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    if not os.path.exists(libreoffice_path):
        # Fallback to standard linux/brew path if they symlinked it
        libreoffice_path = "libreoffice"
        
    env_dir = f"file:///tmp/LibreOffice_Conversion_{uuid.uuid4().hex}"
        
    cmd = [
        libreoffice_path,
        f"-env:UserInstallation={env_dir}",
        "--headless",
        "--convert-to",
        "pdf",
        "--outdir",
        output_dir,
        input_path
    ]
    subprocess.run(cmd, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)

@router.post("/word-to-pdf")
async def word_to_pdf(file: UploadFile = File(...)):
    with tempfile.TemporaryDirectory() as tmpdir:
        input_path = os.path.join(tmpdir, file.filename)
        with open(input_path, "wb") as f:
            f.write(await file.read())
            
        try:
            convert_with_libreoffice(input_path, tmpdir)
            
            # The output filename will be the same but with .pdf extension
            base_name = os.path.splitext(file.filename)[0]
            output_path = os.path.join(tmpdir, f"{base_name}.pdf")
            
            if os.path.exists(output_path):
                # Read into memory so we can return it after tempdir is destroyed
                with open(output_path, "rb") as f:
                    pdf_data = f.read()
                return StreamingResponse(BytesIO(pdf_data), media_type="application/pdf", headers={"Content-Disposition": f"attachment; filename={base_name}.pdf"})
            else:
                raise HTTPException(status_code=500, detail="Conversion failed: Output file not found.")
        except subprocess.CalledProcessError as e:
            raise HTTPException(status_code=500, detail=f"LibreOffice error: {e.stderr.decode()}")

from pdf2docx import Converter

@router.post("/pdf-to-word")
async def pdf_to_word(file: UploadFile = File(...)):
    with tempfile.TemporaryDirectory() as tmpdir:
        input_path = os.path.join(tmpdir, file.filename)
        base_name = os.path.splitext(file.filename)[0]
        output_path = os.path.join(tmpdir, f"{base_name}.docx")
        
        with open(input_path, "wb") as f:
            f.write(await file.read())
            
        try:
            cv = Converter(input_path)
            cv.convert(output_path, start=0, end=None)
            cv.close()
            
            with open(output_path, "rb") as f:
                docx_data = f.read()
            return StreamingResponse(BytesIO(docx_data), media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={"Content-Disposition": f"attachment; filename={base_name}.docx"})
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"pdf2docx error: {str(e)}")

@router.post("/excel-to-pdf")
async def excel_to_pdf(file: UploadFile = File(...)):
    with tempfile.TemporaryDirectory() as tmpdir:
        input_path = os.path.join(tmpdir, file.filename)
        with open(input_path, "wb") as f:
            f.write(await file.read())
            
        try:
            convert_with_libreoffice(input_path, tmpdir)
            base_name = os.path.splitext(file.filename)[0]
            output_path = os.path.join(tmpdir, f"{base_name}.pdf")
            
            if os.path.exists(output_path):
                with open(output_path, "rb") as f:
                    pdf_data = f.read()
                return StreamingResponse(BytesIO(pdf_data), media_type="application/pdf", headers={"Content-Disposition": f"attachment; filename={base_name}.pdf"})
            else:
                raise HTTPException(status_code=500, detail="Conversion failed: Output file not found.")
        except subprocess.CalledProcessError as e:
            raise HTTPException(status_code=500, detail=f"LibreOffice error: {e.stderr.decode()}")

@router.post("/ppt-to-pdf")
async def ppt_to_pdf(file: UploadFile = File(...)):
    with tempfile.TemporaryDirectory() as tmpdir:
        input_path = os.path.join(tmpdir, file.filename)
        with open(input_path, "wb") as f:
            f.write(await file.read())
            
        try:
            convert_with_libreoffice(input_path, tmpdir)
            base_name = os.path.splitext(file.filename)[0]
            output_path = os.path.join(tmpdir, f"{base_name}.pdf")
            
            if os.path.exists(output_path):
                with open(output_path, "rb") as f:
                    pdf_data = f.read()
                return StreamingResponse(BytesIO(pdf_data), media_type="application/pdf", headers={"Content-Disposition": f"attachment; filename={base_name}.pdf"})
            else:
                raise HTTPException(status_code=500, detail="Conversion failed: Output file not found.")
        except subprocess.CalledProcessError as e:
            raise HTTPException(status_code=500, detail=f"LibreOffice error: {e.stderr.decode()}")
@router.post("/id-card-gen")
def id_card_gen(req: IDCardRequest):
    prompt = f"Generate JSON data for a professional {req.card_type} ID card. Name: {req.name}, Role: {req.role}, Org: {req.organization}, ID: {req.id_number}, Blood Group: {req.blood_group}. Return a JSON object with keys: 'front_details', 'back_details', 'barcode_value'."
    raw_json = clean_json_string(generate_ai_response(prompt, json_mode=True))
    try:
        return {"id_card": json.loads(raw_json)}
    except json.JSONDecodeError:
        return {"id_card": raw_json}

import zipfile

@router.post("/pdf-to-image")
async def pdf_to_image(file: UploadFile = File(...), output_format: str = Form("jpeg")):
    # Convert PDF pages directly to images using pdf2image
    pdf_bytes = await file.read()
    output_format = output_format.lower()
    
    if output_format not in ["jpeg", "jpg", "png", "webp"]:
        output_format = "jpeg"
        
    pil_format = "JPEG"
    ext = "jpg"
    if output_format == "png":
        pil_format = "PNG"
        ext = "png"
    elif output_format == "webp":
        pil_format = "WEBP"
        ext = "webp"
    
    try:
        pages = convert_from_bytes(pdf_bytes, dpi=200)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to convert PDF to images: {str(e)}")
        
    out_zip = BytesIO()
    with zipfile.ZipFile(out_zip, "w") as zf:
        for i, page_img in enumerate(pages):
            img_bytes = BytesIO()
            save_kwargs = {}
            if pil_format == "JPEG" or pil_format == "WEBP":
                save_kwargs["quality"] = 85
            page_img.save(img_bytes, format=pil_format, **save_kwargs)
            zf.writestr(f"page_{i+1}.{ext}", img_bytes.getvalue())
                
    if len(pages) == 0:
        raise HTTPException(status_code=404, detail="No pages found in this PDF.")
        
    out_zip.seek(0)
    return StreamingResponse(out_zip, media_type="application/zip", headers={"Content-Disposition": "attachment; filename=extracted_images.zip"})

from pdf2image import convert_from_bytes
import tempfile
import os
import rembg
from typing import Optional
from fastapi import Form
from PIL import ImageDraw, ImageFont

@router.post("/watermark-pdf")
async def watermark_pdf(
    file: UploadFile = File(...), 
    watermark_image: Optional[UploadFile] = File(None),
    watermark_text: Optional[str] = Form(None)
):
    # Read files
    pdf_bytes = await file.read()
    
    wm_img = None
    if watermark_image and watermark_image.filename:
        wm_bytes = await watermark_image.read()
        try:
            no_bg_bytes = rembg.remove(wm_bytes)
            wm_img = Image.open(BytesIO(no_bg_bytes)).convert("RGBA")
        except Exception as e:
            wm_img = Image.open(BytesIO(wm_bytes)).convert("RGBA")
    elif watermark_text:
        # Create a text watermark image
        txt_img = Image.new("RGBA", (3000, 3000), (255, 255, 255, 0))
        d_txt = ImageDraw.Draw(txt_img)
        
        font = None
        font_paths = [
            "/System/Library/Fonts/Helvetica.ttc",
            "/Library/Fonts/Arial.ttf",
            "arial.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"
        ]
        for path in font_paths:
            try:
                font = ImageFont.truetype(path, 200)
                break
            except:
                continue
                
        if font is None:
            font = ImageFont.load_default()
                
        # Draw text center
        d_txt.text((1500, 1500), watermark_text, fill=(100, 100, 100, 255), font=font, anchor="mm")
        
        # Crop to the actual text bounding box so it scales properly
        bbox = txt_img.getbbox()
        if bbox:
            txt_img = txt_img.crop(bbox)
            
        # Rotate text diagonally (45 degrees) and expand
        wm_img = txt_img.rotate(45, expand=True)
    else:
        raise HTTPException(status_code=400, detail="Must provide either watermark_image or watermark_text")
    
    # We will use pdf2image to rasterize the PDF, allowing true transparent overlays!
    try:
        pages = convert_from_bytes(pdf_bytes)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Invalid PDF file (make sure you uploaded a .pdf and not a .docx): {str(e)}")
        
    watermarked_pages = []
    
    for page_img in pages:
        page_img = page_img.convert("RGBA")
        page_width, page_height = page_img.size
        
        # Resize watermark to 80% of page width
        wm_ratio = wm_img.width / wm_img.height
        new_width = int(page_width * 0.8)
        new_height = int(new_width / wm_ratio)
        wm_resized = wm_img.resize((new_width, new_height), Image.LANCZOS)
        
        # Apply opacity (e.g. 30% visible)
        alpha = wm_resized.split()[3]
        alpha = alpha.point(lambda p: p * 0.3)
        wm_resized.putalpha(alpha)
        
        # Calculate center
        x = int((page_width - new_width) / 2)
        y = int((page_height - new_height) / 2)
        
        # Paste with true transparency mask!
        page_img.paste(wm_resized, (x, y), wm_resized)
        
        # Convert back to RGB for PDF saving
        watermarked_pages.append(page_img.convert("RGB"))
        
    if not watermarked_pages:
        raise HTTPException(status_code=400, detail="PDF has no pages.")
        
    out_stream = BytesIO()
    # Save the first page, append the rest
    watermarked_pages[0].save(out_stream, format="PDF", save_all=True, append_images=watermarked_pages[1:])
    out_stream.seek(0)
    
    return StreamingResponse(out_stream, media_type="application/pdf", headers={"Content-Disposition": "attachment; filename=watermarked.pdf"})

@router.post("/pdf-to-text")
async def pdf_to_text(file: UploadFile = File(...)):
    try:
        reader = PdfReader(BytesIO(await file.read()))
        text = "\n".join([page.extract_text() or "" for page in reader.pages])
        return {"extracted_text": text.strip()}
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to read PDF: {str(e)}")

import docx

@router.post("/word-to-text")
async def word_to_text(file: UploadFile = File(...)):
    try:
        doc = docx.Document(BytesIO(await file.read()))
        text = "\n".join([p.text for p in doc.paragraphs])
        return {"extracted_text": text.strip()}
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to read Word document: {str(e)}")

@router.post("/digital-sign")
async def digital_sign(file: UploadFile = File(...), signature_image: UploadFile = File(...), position: str = Form("bottom_right")):
    # Very similar to watermark, but we place it at the selected position
    pdf_bytes = await file.read()
    reader = PdfReader(BytesIO(pdf_bytes))
    writer = PdfWriter()
    
    sig_bytes = await signature_image.read()
    sig_img = Image.open(BytesIO(sig_bytes)).convert("RGBA")
    
    # We only sign the LAST page
    last_page_idx = len(reader.pages) - 1
    
    for i, page in enumerate(reader.pages):
        if i == last_page_idx:
            page_width = float(page.mediabox.width)
            page_height = float(page.mediabox.height)
            
            # Make signature 25% of page width
            sig_ratio = sig_img.width / sig_img.height
            new_width = int(page_width * 0.25)
            new_height = int(new_width / sig_ratio)
            sig_resized = sig_img.resize((new_width, new_height), Image.LANCZOS)
            
            padding = 40
            # If we save RGBA as PDF, PIL might complain or drop transparency. Let's create a white page with the signature, 
            # wait, if we merge a white page, it will hide the PDF content!
            # We must use transparency. PIL saving to PDF doesn't support transparency masks easily for merging in pypdf.
            # Let's just create a small PDF of exactly the signature size, and use PyPDF's merge_translated_page!
            
            # Save the signature itself as a PDF
            sig_pdf_bytes = BytesIO()
            # Convert to RGB with white background (transparency lost, but typically signatures are black on white)
            bg = Image.new("RGB", sig_resized.size, (255, 255, 255))
            bg.paste(sig_resized, mask=sig_resized.split()[3]) # Use alpha channel as mask
            bg.save(sig_pdf_bytes, format="PDF")
            sig_pdf_bytes.seek(0)
            
            sig_reader = PdfReader(sig_pdf_bytes)
            
            # In PDF coordinates, (0,0) is bottom-left
            if position == "bottom_right":
                pdf_x = float(page_width - new_width - padding)
                pdf_y = float(padding)
            elif position == "bottom_left":
                pdf_x = float(padding)
                pdf_y = float(padding)
            elif position == "top_right":
                pdf_x = float(page_width - new_width - padding)
                pdf_y = float(page_height - new_height - padding)
            elif position == "top_left":
                pdf_x = float(padding)
                pdf_y = float(page_height - new_height - padding)
            elif position == "center":
                pdf_x = float((page_width - new_width) / 2)
                pdf_y = float((page_height - new_height) / 2)
            else:
                pdf_x = float(page_width - new_width - padding)
                pdf_y = float(padding) 
            
            # We need to translate the signature page and merge it
            page.merge_translated_page(sig_reader.pages[0], tx=pdf_x, ty=pdf_y)
            writer.add_page(page)
        else:
            writer.add_page(page)
            
    out_stream = BytesIO()
    writer.write(out_stream)
    out_stream.seek(0)
    return StreamingResponse(out_stream, media_type="application/pdf", headers={"Content-Disposition": "attachment; filename=signed.pdf"})

@router.post("/pdf-resize")
async def pdf_resize(file: UploadFile = File(...), size: str = Form("A4")):
    sizes = {
        "A1": (1684, 2384),
        "A2": (1191, 1684),
        "A3": (842, 1191),
        "A4": (595.28, 841.89),
        "A5": (420, 595),
        "B3": (1001, 1417),
        "B4": (709, 1001),
        "B5": (499, 709),
        "LETTER": (612, 792),
        "NOTE": (540, 720),
        "LEGAL": (612, 1008),
        "TABLOID": (792, 1224),
        "EXECUTIVE": (522, 756),
        "POSTCARD": (283, 416)
    }
    
    target_size = sizes.get(size.upper())
    if not target_size:
        raise HTTPException(status_code=400, detail=f"Invalid size: {size}")
        
    reader = PdfReader(BytesIO(await file.read()))
    writer = PdfWriter()
    
    for page in reader.pages:
        page.scale_to(width=target_size[0], height=target_size[1])
        writer.add_page(page)
        
    out_stream = BytesIO()
    writer.write(out_stream)
    out_stream.seek(0)
    return StreamingResponse(out_stream, media_type="application/pdf", headers={"Content-Disposition": "attachment; filename=resized.pdf"})

@router.post("/pdf-organize")
async def pdf_organize(file: UploadFile = File(...), rotations: str = Form("{}"), deletions: str = Form("")):
    import json
    try:
        rot_dict = json.loads(rotations)
    except:
        rot_dict = {}
        
    del_set = set([int(x.strip()) for x in deletions.split(",") if x.strip().isdigit()])
    
    reader = PdfReader(BytesIO(await file.read()))
    writer = PdfWriter()
    
    for i, page in enumerate(reader.pages):
        if i in del_set:
            continue
            
        str_i = str(i)
        if str_i in rot_dict:
            page.rotate(int(rot_dict[str_i]))
            
        writer.add_page(page)
        
    out_stream = BytesIO()
    writer.write(out_stream)
    out_stream.seek(0)
    return StreamingResponse(out_stream, media_type="application/pdf", headers={"Content-Disposition": "attachment; filename=organized.pdf"})

@router.post("/pdf-add-images")
async def pdf_add_images(pdf_file: UploadFile = File(...), images: List[UploadFile] = File(...)):
    reader = PdfReader(BytesIO(await pdf_file.read()))
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
        
    for img_file in images:
        img_bytes = await img_file.read()
        img = Image.open(BytesIO(img_bytes)).convert("RGB")
        img_pdf = BytesIO()
        img.save(img_pdf, format="PDF")
        img_pdf.seek(0)
        img_reader = PdfReader(img_pdf)
        writer.add_page(img_reader.pages[0])
        
    out_stream = BytesIO()
    writer.write(out_stream)
    out_stream.seek(0)
    return StreamingResponse(out_stream, media_type="application/pdf", headers={"Content-Disposition": "attachment; filename=appended.pdf"})

@router.post("/pdf-protect")
async def pdf_protect(
    file: UploadFile = File(...), 
    password: str = Form(""), # Legacy fallback
    user_password: str = Form(""),
    owner_password: str = Form("")
):
    actual_user_pwd = user_password or password
    
    if not actual_user_pwd and not owner_password:
        raise HTTPException(status_code=400, detail="Must provide at least one password to lock the PDF.")
        
    actual_owner_pwd = owner_password or actual_user_pwd

    reader = PdfReader(BytesIO(await file.read()))
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
        
    writer.encrypt(user_password=actual_user_pwd, owner_password=actual_owner_pwd)
    out_stream = BytesIO()
    writer.write(out_stream)
    out_stream.seek(0)
    return StreamingResponse(out_stream, media_type="application/pdf", headers={"Content-Disposition": "attachment; filename=protected.pdf"})

@router.post("/pdf-unlock")
async def pdf_unlock(file: UploadFile = File(...), password: str = Form("")):
    reader = PdfReader(BytesIO(await file.read()))
    if not reader.is_encrypted:
        raise HTTPException(status_code=400, detail="PDF is not encrypted.")
        
    # First try an empty password to automatically bypass Owner/Permissions locks
    if not reader.decrypt(""):
        # If empty fails, a User Password is required. Try the provided password.
        if not password or not reader.decrypt(password):
            raise HTTPException(status_code=401, detail="Incorrect password. A user password is required to open this file.")

    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
        
    out_stream = BytesIO()
    writer.write(out_stream)
    out_stream.seek(0)
    return StreamingResponse(out_stream, media_type="application/pdf", headers={"Content-Disposition": "attachment; filename=unlocked.pdf"})

from fpdf import FPDF

@router.post("/text-to-pdf")
async def text_to_pdf(text_content: str = Form(...)):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.multi_cell(0, 10, txt=text_content)
    
    out_bytes = pdf.output(dest='S')
    if isinstance(out_bytes, str):
        out_bytes = out_bytes.encode('latin-1', 'replace')
        
    return StreamingResponse(BytesIO(out_bytes), media_type="application/pdf", headers={"Content-Disposition": "attachment; filename=text.pdf"})

import cv2
import numpy as np

@router.post("/detect-edges")
async def detect_edges(image: UploadFile = File(...)):
    img_bytes = await image.read()
    nparr = np.frombuffer(img_bytes, np.uint8)
    img = cv2.imdecode(nparr, cv2.IMREAD_COLOR)
    
    if img is None:
        raise HTTPException(status_code=400, detail="Invalid image file.")
        
    ratio = img.shape[0] / 500.0
    orig = img.copy()
    image_resized = cv2.resize(img, (int(img.shape[1] / ratio), 500))
    
    gray = cv2.cvtColor(image_resized, cv2.COLOR_BGR2GRAY)
    gray = cv2.GaussianBlur(gray, (5, 5), 0)
    edged = cv2.Canny(gray, 75, 200)
    
    cnts, _ = cv2.findContours(edged.copy(), cv2.RETR_LIST, cv2.CHAIN_APPROX_SIMPLE)
    cnts = sorted(cnts, key=cv2.contourArea, reverse=True)[:5]
    
    screenCnt = None
    for c in cnts:
        peri = cv2.arcLength(c, True)
        approx = cv2.approxPolyDP(c, 0.02 * peri, True)
        if len(approx) == 4:
            screenCnt = approx
            break
            
    if screenCnt is None:
        h, w = orig.shape[:2]
        corners = [{"x": 0, "y": 0}, {"x": w, "y": 0}, {"x": w, "y": h}, {"x": 0, "y": h}]
    else:
        screenCnt = (screenCnt.reshape(4, 2) * ratio).astype(int)
        
        def order_points(pts):
            rect = np.zeros((4, 2), dtype="float32")
            s = pts.sum(axis=1)
            rect[0] = pts[np.argmin(s)]
            rect[2] = pts[np.argmax(s)]
            diff = np.diff(pts, axis=1)
            rect[1] = pts[np.argmin(diff)]
            rect[3] = pts[np.argmax(diff)]
            return rect
            
        rect = order_points(screenCnt)
        corners = [
            {"x": int(rect[0][0]), "y": int(rect[0][1])},
            {"x": int(rect[1][0]), "y": int(rect[1][1])},
            {"x": int(rect[2][0]), "y": int(rect[2][1])},
            {"x": int(rect[3][0]), "y": int(rect[3][1])},
        ]
        
    return {"corners": corners, "width": img.shape[1], "height": img.shape[0]}

@router.post("/process-scan-cropped")
async def process_scan_cropped(
    image: UploadFile = File(...),
    corners: str = Form(...),
    scan_type: str = Form("magic_color")
):
    import json
    img_bytes = await image.read()
    nparr = np.frombuffer(img_bytes, np.uint8)
    img = cv2.imdecode(nparr, cv2.IMREAD_COLOR)
    
    try:
        pts = json.loads(corners)
        pts_np = np.array([[p["x"], p["y"]] for p in pts], dtype="float32")
    except:
        raise HTTPException(status_code=400, detail="Invalid corners format")
        
    (tl, tr, br, bl) = pts_np
    widthA = np.sqrt(((br[0] - bl[0]) ** 2) + ((br[1] - bl[1]) ** 2))
    widthB = np.sqrt(((tr[0] - tl[0]) ** 2) + ((tr[1] - tl[1]) ** 2))
    maxWidth = max(int(widthA), int(widthB))
    
    heightA = np.sqrt(((tr[0] - br[0]) ** 2) + ((tr[1] - br[1]) ** 2))
    heightB = np.sqrt(((tl[0] - bl[0]) ** 2) + ((tl[1] - bl[1]) ** 2))
    maxHeight = max(int(heightA), int(heightB))
    
    dst = np.array([
        [0, 0],
        [maxWidth - 1, 0],
        [maxWidth - 1, maxHeight - 1],
        [0, maxHeight - 1]], dtype="float32")
        
    M = cv2.getPerspectiveTransform(pts_np, dst)
    warped = cv2.warpPerspective(img, M, (maxWidth, maxHeight))
    
    warped_rgb = cv2.cvtColor(warped, cv2.COLOR_BGR2RGB)
    pil_img = Image.fromarray(warped_rgb)
    
    if scan_type == "magic_color":
        enhancer = ImageEnhance.Contrast(pil_img)
        pil_img = enhancer.enhance(1.8)
        enhancer = ImageEnhance.Brightness(pil_img)
        pil_img = enhancer.enhance(1.1)
        enhancer = ImageEnhance.Color(pil_img)
        pil_img = enhancer.enhance(1.5)
    elif scan_type == "black_white":
        pil_img = pil_img.convert("L")
        pil_img = pil_img.point(lambda x: 0 if x < 128 else 255, '1')
        pil_img = pil_img.convert("RGB")
    elif scan_type == "grayscale":
        pil_img = pil_img.convert("L").convert("RGB")
        
    out_stream = BytesIO()
    pil_img.save(out_stream, format="JPEG", quality=90)
    out_stream.seek(0)
    return StreamingResponse(out_stream, media_type="image/jpeg", headers={"Content-Disposition": "attachment; filename=cropped_scan.jpg"})


@router.post("/unlock-pdf")
async def unlock_pdf(file: UploadFile = File(...), password: str = Form("")):
    try:
        reader = PdfReader(BytesIO(await file.read()))
        if reader.is_encrypted:
            reader.decrypt(password)
        writer = PdfWriter()
        for page in reader.pages:
            writer.add_page(page)
        out_stream = BytesIO()
        writer.write(out_stream)
        out_stream.seek(0)
        return StreamingResponse(out_stream, media_type="application/pdf", headers={"Content-Disposition": "attachment; filename=unlocked.pdf"})
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to unlock PDF: {str(e)}")

@router.post("/pdf-to-text")
async def pdf_to_text(file: UploadFile = File(...)):
    try:
        reader = PdfReader(BytesIO(await file.read()))
        text = ""
        for page in reader.pages:
            extracted = page.extract_text()
            if extracted:
                text += extracted + "\n"
        return {"extracted_text": text}
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to extract text: {str(e)}")

from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from reportlab.lib.utils import simpleSplit

@router.post("/text-to-pdf")
async def text_to_pdf(text: str = Form(...)):
    try:
        out_stream = BytesIO()
        c = canvas.Canvas(out_stream, pagesize=letter)
        width, height = letter
        
        lines = text.split('\n')
        y = height - 40
        for line in lines:
            wrapped = simpleSplit(line, 'Helvetica', 12, width - 80)
            for w_line in wrapped:
                if y < 40:
                    c.showPage()
                    y = height - 40
                c.drawString(40, y, w_line)
                y -= 15
        c.save()
        out_stream.seek(0)
        return StreamingResponse(out_stream, media_type="application/pdf", headers={"Content-Disposition": "attachment; filename=document.pdf"})
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to create PDF: {str(e)}")

from reportlab.lib.pagesizes import A1, A2, A3, A4, A5, B3, B4, B5, letter, legal

PAGE_SIZES = {
    "A1": A1,
    "A2": A2,
    "A3": A3,
    "A4": A4,
    "A5": A5,
    "B3": B3,
    "B4": B4,
    "B5": B5,
    "LETTER": letter,
    "LEGAL": legal,
    "TABLOID": (11 * 72.0, 17 * 72.0)
}

@router.post("/modify-pages")
async def modify_pages(
    file: UploadFile = File(...),
    delete_pages: str = Form(""),
    rotate_pages: str = Form(""),
    rotation_angle: int = Form(90),
    page_size: str = Form("auto")
):
    try:
        reader = PdfReader(BytesIO(await file.read()))
        writer = PdfWriter()
        
        def parse_pages(page_str, max_pages):
            pages = set()
            for part in page_str.split(','):
                part = part.strip()
                if not part: continue
                if '-' in part:
                    try:
                        start, end = part.split('-')
                        if start.isdigit() and end.isdigit():
                            pages.update(range(int(start) - 1, int(end)))
                    except: pass
                elif part.isdigit():
                    pages.add(int(part) - 1)
            return {p for p in pages if 0 <= p < max_pages}

        to_delete = parse_pages(delete_pages, len(reader.pages))
        to_rotate = parse_pages(rotate_pages, len(reader.pages))
        
        for i, page in enumerate(reader.pages):
            if i in to_delete:
                continue
            if i in to_rotate:
                page.transfer_rotation_to_content()
                page.rotate(rotation_angle)
            
            if page_size != "auto" and page_size.upper() in PAGE_SIZES:
                target_w, target_h = PAGE_SIZES[page_size.upper()]
                page.scale_to(target_w, target_h)
                
            writer.add_page(page)
            
        out_stream = BytesIO()
        writer.write(out_stream)
        out_stream.seek(0)
        return StreamingResponse(out_stream, media_type="application/pdf", headers={"Content-Disposition": "attachment; filename=modified.pdf"})
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to modify PDF: {str(e)}")

@router.post("/pdf-to-excel")
async def pdf_to_excel(file: UploadFile = File(...)):
    if not file.filename.lower().endswith('.pdf'):
        raise HTTPException(status_code=400, detail="File must be a PDF")
        
    import io
    import pandas as pd
    import pdfplumber
    
    try:
        pdf_bytes = await file.read()
        excel_buffer = io.BytesIO()
        
        with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
            with pd.ExcelWriter(excel_buffer, engine='openpyxl') as writer:
                table_found = False
                for i, page in enumerate(pdf.pages):
                    tables = page.extract_tables()
                    
                    # Fallback for borderless tables
                    if not tables:
                        tables = page.extract_tables(table_settings={
                            "vertical_strategy": "text", 
                            "horizontal_strategy": "text"
                        })
                        
                    for j, table in enumerate(tables):
                        # Filter out empty rows or None values
                        clean_table = [[cell for cell in row] for row in table if any(cell for cell in row)]
                        if len(clean_table) > 1:
                            df = pd.DataFrame(clean_table[1:], columns=clean_table[0] if clean_table[0] else None)
                            df.to_excel(writer, sheet_name=f'Page_{i+1}_Table_{j+1}', index=False)
                            table_found = True
                        elif len(clean_table) == 1:
                            df = pd.DataFrame(clean_table)
                            df.to_excel(writer, sheet_name=f'Page_{i+1}_Table_{j+1}', index=False, header=False)
                            table_found = True
                            
                if not table_found:
                    for i, page in enumerate(pdf.pages):
                        text = page.extract_text()
                        if text:
                            lines = [line.split() for line in text.split("\n")]
                            df = pd.DataFrame(lines)
                            df.to_excel(writer, sheet_name=f"Page_{i+1}_Text", index=False, header=False)
                            table_found = True
                            
                if not table_found:
                    pd.DataFrame([["No data detected in PDF"]]).to_excel(writer, sheet_name="Result", index=False, header=False)
                    
        excel_buffer.seek(0)
        
        return StreamingResponse(
            excel_buffer,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={file.filename.rsplit('.', 1)[0]}.xlsx"}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"PDF to Excel failed: {str(e)}")

@router.post("/pdf-to-ppt")
async def pdf_to_ppt(file: UploadFile = File(...)):
    if not file.filename.lower().endswith('.pdf'):
        raise HTTPException(status_code=400, detail="File must be a PDF")
        
    import io
    from pdf2image import convert_from_bytes
    from pptx import Presentation
    from pptx.util import Inches
    
    try:
        pdf_bytes = await file.read()
        
        # Convert PDF to images
        images = convert_from_bytes(pdf_bytes)
        
        prs = Presentation()
        # Use blank slide layout
        blank_slide_layout = prs.slide_layouts[6]
        
        for img in images:
            slide = prs.slides.add_slide(blank_slide_layout)
            
            img_buffer = io.BytesIO()
            img.save(img_buffer, format="PNG")
            img_buffer.seek(0)
            
            left = top = Inches(0)
            # Add image to slide, fitting the slide dimensions
            slide.shapes.add_picture(img_buffer, left, top, width=prs.slide_width, height=prs.slide_height)
            
        ppt_buffer = io.BytesIO()
        prs.save(ppt_buffer)
        ppt_buffer.seek(0)
        
        return StreamingResponse(
            ppt_buffer,
            media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            headers={"Content-Disposition": f"attachment; filename={file.filename.rsplit('.', 1)[0]}.pptx"}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"PDF to PPT failed: {str(e)}")


@router.post("/excel-to-csv")
async def excel_to_csv(file: UploadFile = File(...)):
    if not file.filename.lower().endswith((".xls", ".xlsx")):
        raise HTTPException(status_code=400, detail="File must be an Excel file")
    import pandas as pd
    import io
    try:
        content = await file.read()
        df = pd.read_excel(io.BytesIO(content))
        csv_buffer = io.StringIO()
        df.to_csv(csv_buffer, index=False)
        csv_buffer.seek(0)
        return StreamingResponse(
            iter([csv_buffer.getvalue()]),
            media_type="text/csv",
            headers={"Content-Disposition": f"attachment; filename={file.filename.rsplit(".", 1)[0]}.csv"}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@router.post("/csv-to-excel")
async def csv_to_excel(file: UploadFile = File(...)):
    if not file.filename.lower().endswith(".csv"):
        raise HTTPException(status_code=400, detail="File must be a CSV file")
    import pandas as pd
    import io
    try:
        content = await file.read()
        df = pd.read_csv(io.BytesIO(content))
        excel_buffer = io.BytesIO()
        df.to_excel(excel_buffer, index=False)
        excel_buffer.seek(0)
        return StreamingResponse(
            excel_buffer,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={file.filename.rsplit(".", 1)[0]}.xlsx"}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@router.post("/csv-to-pdf")
async def csv_to_pdf(file: UploadFile = File(...)):
    if not file.filename.lower().endswith(".csv"):
        raise HTTPException(status_code=400, detail="File must be a CSV file")
    import pandas as pd
    import io
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
    from reportlab.lib import colors

    try:
        content = await file.read()
        df = pd.read_csv(io.BytesIO(content))
        
        # Convert DataFrame to a list of lists, including the header row
        data = [df.columns.values.astype(str).tolist()] + df.astype(str).values.tolist()
        
        pdf_buffer = io.BytesIO()
        # Use landscape to better fit wide CSV tables
        doc = SimpleDocTemplate(pdf_buffer, pagesize=landscape(letter), leftMargin=30, rightMargin=30, topMargin=30, bottomMargin=30)
        
        t = Table(data)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#333333")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
            ('TOPPADDING', (0, 0), (-1, 0), 8),
            ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor("#f8f9fa")),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey)
        ]))
        
        doc.build([t])
        pdf_buffer.seek(0)
        
        return StreamingResponse(
            pdf_buffer,
            media_type="application/pdf",
            headers={"Content-Disposition": f"attachment; filename={file.filename.rsplit('.', 1)[0]}.pdf"}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"PDF generation failed: {str(e)}")
